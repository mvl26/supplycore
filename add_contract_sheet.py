# -*- coding: utf-8 -*-
"""Thêm sheet Hợp đồng khung (Framework Contract) vào file master data, cùng format + dòng fieldname."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PATH = "/home/hoangvietyeuem/frappe-bench/apps/supplycore/docs/SupplyCore_Master_Data_Import.xlsx"
wb = openpyxl.load_workbook(PATH)

F = "Calibri"; SZ = 14
fn_font = Font(name="Consolas", size=10, bold=True, color="FF7030A0")
h_font = Font(name=F, size=SZ, bold=True, color="FF1A237E")
d_font = Font(name=F, size=SZ, bold=False, color="FF555555")
s_font = Font(name=F, size=SZ, bold=False, color="FF1B5E20")
e_font = Font(name=F, size=SZ, color="FF000000")
fn_fill = PatternFill("solid", fgColor="FFF2E6FF")
h_fill = PatternFill("solid", fgColor="FFFFF9C4")
d_fill = PatternFill("solid", fgColor="FFD6EAF8")
s_fill = PatternFill("solid", fgColor="FFEAFAF1")
w_fill = PatternFill("solid", fgColor="FFFFFFFF")
thin = Side(style="thin", color="FFD0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
lefttop = Alignment(horizontal="left", vertical="top", wrap_text=True)
leftmid = Alignment(horizontal="left", vertical="center", wrap_text=True)

# (fieldname, label, required, description, sample1, sample2)
# Cha = thông tin hợp đồng (lặp/để trống ở dòng vật tư thứ 2+). Con = items.* (mỗi dòng 1 vật tư).
FIELDS = [
    ("supplier", "Nhà cung cấp (*)", True,
     "Mã NCC (khớp sheet Nhà cung cấp, dùng mã SC-SUP-#####). Chỉ điền ở DÒNG ĐẦU của mỗi hợp đồng.",
     "SC-SUP-00001", ""),
    ("contract_number", "Số hợp đồng (*)", True,
     "Số hợp đồng duy nhất. VD: HD-2026-001. Chỉ điền ở DÒNG ĐẦU của mỗi hợp đồng.",
     "HD-2026-001", ""),
    ("contract_date", "Ngày ký (*)", True,
     "Ngày ký hợp đồng. Định dạng YYYY-MM-DD. Chỉ điền ở dòng đầu.",
     "2026-01-05", ""),
    ("valid_from", "Hiệu lực từ (*)", True,
     "Ngày bắt đầu hiệu lực. Định dạng YYYY-MM-DD. Chỉ điền ở dòng đầu.",
     "2026-01-10", ""),
    ("valid_to", "Hết hạn (*)", True,
     "Ngày hết hạn hợp đồng. Định dạng YYYY-MM-DD. Chỉ điền ở dòng đầu.",
     "2026-12-31", ""),
    ("payment_terms", "Điều khoản thanh toán", False,
     "VD: Net 30, COD... Chỉ điền ở dòng đầu.",
     "Net 30", ""),
    ("delivery_terms", "Điều khoản giao hàng", False,
     "Mô tả điều khoản giao hàng. Chỉ điền ở dòng đầu.",
     "Giao tại kho bệnh viện trong 7 ngày", ""),
    ("remarks", "Ghi chú", False,
     "Ghi chú thêm về hợp đồng. Chỉ điền ở dòng đầu.",
     "HĐ khung vật tư tiêu hao năm 2026", ""),
    ("items.item_code", "Mã VT (*)", True,
     "Mã vật tư trong hợp đồng (khớp sheet Vật tư). Mỗi DÒNG là một vật tư.",
     "VT-00001", "VT-00002"),
    ("items.uom", "Đơn vị (*)", True,
     "Đơn vị tính của vật tư (khớp sheet Đơn vị tính).",
     "Cái", "Hộp"),
    ("items.contract_qty", "SL hợp đồng (*)", True,
     "Số lượng cam kết theo hợp đồng cho vật tư này.",
     "10000", "500"),
    ("items.unit_price", "Đơn giá (VND) (*)", True,
     "Đơn giá vật tư theo hợp đồng (VND, không dấu chấm/phẩy).",
     "2500", "120000"),
]

EMPTY_ROWS = 50
title = "Hợp đồng khung (Framework Contract)"[:31]
if title in wb.sheetnames:
    del wb[title]
ws = wb.create_sheet(title=title)
ws.freeze_panes = "A6"   # giữ fieldname + nhãn + mô tả + 2 dòng mẫu

for i, (fieldname, label, reqd, desc, s1, s2) in enumerate(FIELDS, start=1):
    col = get_column_letter(i)
    c = ws.cell(row=1, column=i, value=fieldname); c.font = fn_font; c.fill = fn_fill; c.alignment = center; c.border = border
    c = ws.cell(row=2, column=i, value=label);     c.font = h_font;  c.fill = h_fill;  c.alignment = center; c.border = border
    c = ws.cell(row=3, column=i, value=desc);       c.font = d_font;  c.fill = d_fill;  c.alignment = lefttop; c.border = border
    c = ws.cell(row=4, column=i, value=s1);         c.font = s_font;  c.fill = s_fill;  c.alignment = leftmid; c.border = border
    c = ws.cell(row=5, column=i, value=s2);         c.font = s_font;  c.fill = s_fill;  c.alignment = leftmid; c.border = border
    for r in range(6, 6 + EMPTY_ROWS):
        c = ws.cell(row=r, column=i, value=None); c.font = e_font; c.fill = w_fill; c.alignment = leftmid; c.border = border
    ws.column_dimensions[col].width = min(max(len(fieldname) + 2, len(label) * 0.95, len(str(s1)) + 2, 13), 30)

ws.row_dimensions[1].height = 18
ws.row_dimensions[2].height = 40
ws.row_dimensions[3].height = 70
ws.row_dimensions[4].height = 26
ws.row_dimensions[5].height = 26

# Đặt sheet hợp đồng ngay sau "Vật tư (Item)"
names = wb.sheetnames
ws_obj = wb[title]
wb._sheets.remove(ws_obj)
item_name = next((n for n in names if n.startswith("Vật tư")), None)
idx = wb.sheetnames.index(item_name) + 1 if item_name else len(wb._sheets)
wb._sheets.insert(idx, ws_obj)

wb.save(PATH)
print("Saved. Sheets:", wb.sheetnames)
