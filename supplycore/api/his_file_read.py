"""Đọc file bàn giao từ tool his-slip-extractor → dict canonical cho _process_extracted.
Hỗ trợ JSON (máy) và Excel (người sửa tay — Excel là nguồn sự thật trên đường đó)."""
import json
import frappe
from frappe import _
from frappe.utils import flt, cint

SUPPORTED_VERSION = 1
LINE_COLUMNS = ["tt", "name", "his_code", "uom", "batch_no", "expiry",
                "qty", "unit_price", "amount"]
NUM_COLS = {"tt", "qty", "unit_price", "amount"}


def _check_version(v):
    if cint(v) != SUPPORTED_VERSION:
        frappe.throw(_("SC-E-HIS-EXTRACT: schema_version {0} không hỗ trợ (cần {1}). "
                       "Cập nhật tool/SupplyCore.").format(v, SUPPORTED_VERSION),
                     title="SC-E-HIS-EXTRACT")


def read_json_file(path: str) -> dict:
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    _check_version(data.get("schema_version"))
    return data


def read_xlsx_file(path: str) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    if "header" not in wb.sheetnames or "lines" not in wb.sheetnames:
        frappe.throw(_("SC-E-HIS-EXTRACT: Excel thiếu sheet 'header'/'lines'"),
                     title="SC-E-HIS-EXTRACT")
    hdr = {r[0].value: r[1].value for r in wb["header"].iter_rows() if r[0].value}
    _check_version(hdr.get("schema_version"))
    ls = wb["lines"]
    cols = [c.value for c in ls[1]]
    if cols[:len(LINE_COLUMNS)] != LINE_COLUMNS:
        frappe.throw(_("SC-E-HIS-EXTRACT: cột sheet 'lines' sai layout hợp đồng"),
                     title="SC-E-HIS-EXTRACT")
    lines = []
    for row in ls.iter_rows(min_row=2):
        vals = {cols[i]: (c.value if c.value is not None else "") for i, c in enumerate(row)}
        if all((v == "" or v is None) for v in vals.values()):
            continue
        line = {}
        for k in LINE_COLUMNS:
            v = vals.get(k, "")
            line[k] = flt(v) if k in NUM_COLS else str(v).strip()
        line["tt"] = cint(line["tt"])
        lines.append(line)
    return {
        "schema_version": SUPPORTED_VERSION,
        "slip_type": hdr.get("slip_type", "transfer"),
        "profile": hdr.get("profile", ""),
        "slip_no": str(hdr.get("slip_no", "")).strip(),
        "slip_date": str(hdr.get("slip_date", "")).strip(),
        "from_warehouse_name": str(hdr.get("from_warehouse_name", "")).strip(),
        "to_warehouse_name": str(hdr.get("to_warehouse_name", "")).strip(),
        "lines": lines,
    }
