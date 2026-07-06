"""SupplyCore — Xuất/Nhập phiếu cha-con bằng Excel 2 sheet (engine cấu hình chung).

Một engine duy nhất cho mọi chứng từ có 1 child table chính:
  - Sheet "<phiếu>" : mỗi phiếu 1 dòng (thông tin chung).
  - Sheet "Vật tư"  : mỗi dòng vật tư 1 dòng, cột Mã phiếu trỏ về phiếu cha.
  - Sheet "Hướng dẫn": chèn trong template, giải thích Mã phiếu + quy tắc.

Nhập: cập nhật phiếu đã có (theo mã) + tạo mới phiếu chưa có. **Phiếu nhập vào LUÔN
lưu ở dạng Draft** (chỉ insert/save, KHÔNG bao giờ submit). Nhập đè danh mục vật tư
CHỈ áp dụng phiếu đang Draft (docstatus=0) — bảo vệ phiếu đã duyệt/đã ghi sổ.

Cấu hình từng doctype trong CONFIGS. API path: supplycore.api.voucher_io.*
Xem flow: supplycore/m1_contract/FC_IMPORT_EXPORT_FLOW.md
"""

from __future__ import annotations

import base64
import io
from datetime import date, datetime
from typing import Any

import frappe
from frappe import _
from frappe.utils import cint, flt, getdate, get_datetime

from supplycore.api.data_io import _check_perm, _safe_filename

SHEET_GUIDE = "Hướng dẫn"
SHEET_ITEM = "Vật tư"

# Định dạng tiền tệ VND cho ô Excel: hiển thị "1.234.567 ₫" nhưng GIÁ TRỊ vẫn là số
# (load_workbook data_only đọc ra số) → nhập lại không vỡ. 0 chữ số thập phân (VND).
VND_NUMFMT = '#,##0" ₫"'

# ---------------------------------------------------------------------------
# CẤU HÌNH TỪNG DOCTYPE
#   header_cols / item_cols : (label hiển thị, fieldname) theo đúng thứ tự cột.
#   *_writable : field người dùng được nhập (loại read-only/auto/display).
#   item_key   : field bắt buộc nhận diện 1 dòng vật tư (item / item_code).
#   required_create : field bắt buộc khi TẠO MỚI.
#   link_header / link_item : {fieldname: doctype} để kiểm tra Link tồn tại.
#   editable_stage_field / editable_stages : khoá sửa theo stage (FC). None = chỉ
#     cần docstatus=0 (Draft) là sửa/đè được.
# ---------------------------------------------------------------------------
CONFIGS: dict[str, dict] = {
    "Framework Contract": {
        "label": "Hợp đồng khung",
        "sheet_parent": "Hợp đồng",
        "header_cols": [
            ("Mã HĐ", "name"), ("Số HĐ", "contract_number"), ("Mã NCC", "supplier"),
            ("Tên NCC", "supplier_name"), ("Ngày ký", "contract_date"),
            ("Hiệu lực từ", "valid_from"), ("Hết hạn", "valid_to"),
            ("Tổng giá trị (VND)", "total_value"), ("ĐK thanh toán", "payment_terms"),
            ("ĐK giao hàng", "delivery_terms"), ("Tệp đính kèm", "attachment"),
            ("Người tạo", "owner"), ("Ghi chú", "remarks"),
        ],
        "header_writable": {"contract_number", "supplier", "contract_date", "valid_from",
                            "valid_to", "payment_terms", "delivery_terms", "attachment", "remarks"},
        "child_field": "items", "child_doctype": "FC Item", "item_key": "item_code",
        "item_cols": [
            ("Mã HĐ", "parent"), ("STT", "idx"), ("Mã VT", "item_code"), ("Tên VT", "item_name"),
            ("UOM", "uom"), ("SL HĐ", "contract_qty"), ("Đơn giá (VND)", "unit_price"),
            ("Thành tiền (VND)", "total_amount"),
        ],
        "item_writable": {"item_code", "uom", "contract_qty", "unit_price"},
        "required_create": ["contract_number", "supplier", "contract_date", "valid_from", "valid_to"],
        "link_header": {"supplier": "SC Supplier"},
        "link_item": {"item_code": "SC Item", "uom": "SC UOM"},
        "editable_stage_field": "approval_stage", "editable_stages": {"", "Draft", "Rejected"},
    },
    "SC Transfer Request": {
        "label": "Yêu cầu chuyển kho",
        "sheet_parent": "Phiếu",
        "header_cols": [
            ("Mã phiếu", "name"), ("Ngày yêu cầu", "request_date"), ("Loại chuyển kho", "transfer_type"),
            ("Ngày cần", "required_by"), ("Kho nguồn", "from_warehouse"), ("Kho đích", "to_warehouse"),
            ("Khoa yêu cầu", "requested_for_department"), ("Trạng thái", "status"),
            ("Tổng SL", "total_qty"), ("Người tạo", "owner"), ("Ghi chú", "remarks"),
        ],
        "header_writable": {"request_date", "transfer_type", "required_by", "from_warehouse",
                            "to_warehouse", "requested_for_department", "remarks"},
        "child_field": "items", "child_doctype": "SC Transfer Request Item", "item_key": "item",
        "item_cols": [
            ("Mã phiếu", "parent"), ("STT", "idx"), ("Mã VT", "item"), ("Tên VT", "item_name"),
            ("UOM", "uom"), ("SL yêu cầu", "requested_qty"), ("SL duyệt", "approved_qty"),
            ("Lô", "batch"), ("Ghi chú", "remarks"),
        ],
        "item_writable": {"item", "uom", "requested_qty", "approved_qty", "batch", "remarks"},
        "required_create": ["request_date", "transfer_type", "required_by", "from_warehouse", "to_warehouse"],
        "link_header": {"from_warehouse": "SC Warehouse", "to_warehouse": "SC Warehouse",
                        "requested_for_department": "SC Department"},
        "link_item": {"item": "SC Item", "uom": "SC UOM", "batch": "SC Batch"},
        "editable_stage_field": None, "editable_stages": None,
    },
    "SC Purchase Receipt": {
        "label": "Phiếu nhập (tiếp nhận)",
        "sheet_parent": "Phiếu",
        "header_cols": [
            ("Mã phiếu", "name"), ("NCC", "supplier"), ("Tên NCC", "supplier_name"),
            ("PO tham chiếu", "purchase_order"), ("Ngày nhập", "posting_date"),
            ("Kho nhập", "to_warehouse"), ("Yêu cầu QC", "qc_required"),
            ("Lý do không có PO", "no_po_reason"), ("Trạng thái QC", "qc_status"),
            ("Tổng SL", "total_qty"), ("Tổng giá trị (VND)", "total_value"),
            ("Người tạo", "owner"), ("Ghi chú", "remarks"),
        ],
        "header_writable": {"supplier", "purchase_order", "posting_date", "to_warehouse",
                            "qc_required", "no_po_reason", "remarks"},
        "child_field": "items", "child_doctype": "SC Purchase Receipt Item", "item_key": "item",
        "item_cols": [
            ("Mã phiếu", "parent"), ("STT", "idx"), ("Mã VT", "item"), ("Tên VT", "item_name"),
            ("SL nhận", "qty"), ("UOM", "uom"), ("Đơn giá (VND)", "rate"), ("Thành tiền (VND)", "amount"),
            ("Số lô NCC", "supplier_batch_no"), ("Ngày SX", "manufacturing_date"),
            ("Hạn dùng", "expiry_date"), ("Lô (đã có)", "batch_no"), ("Kho", "warehouse"),
        ],
        "item_writable": {"item", "qty", "uom", "rate", "supplier_batch_no", "manufacturing_date",
                          "expiry_date", "batch_no", "warehouse"},
        "required_create": ["supplier", "posting_date", "to_warehouse"],
        "link_header": {"supplier": "SC Supplier", "to_warehouse": "SC Warehouse",
                        "purchase_order": "SC Purchase Order"},
        "link_item": {"item": "SC Item", "uom": "SC UOM", "batch_no": "SC Batch", "warehouse": "SC Warehouse"},
        "editable_stage_field": None, "editable_stages": None,
    },
    "SC Material Request": {
        "label": "Yêu cầu mua",
        "sheet_parent": "Phiếu",
        "header_cols": [
            ("Mã phiếu", "name"), ("Loại yêu cầu", "request_type"), ("Ngày yêu cầu", "transaction_date"),
            ("Ngày cần", "schedule_date"), ("Khoa phòng", "department"), ("Kho đích", "warehouse"),
            ("Trạng thái", "status"), ("Tổng SL", "total_qty"), ("Tổng ước tính (VND)", "total_estimated_cost"),
            ("Người tạo", "owner"), ("Lý do", "reason"), ("Ghi chú", "remarks"),
        ],
        "header_writable": {"request_type", "transaction_date", "schedule_date", "department",
                            "warehouse", "reason", "remarks"},
        "child_field": "items", "child_doctype": "SC Material Request Item", "item_key": "item",
        "item_cols": [
            ("Mã phiếu", "parent"), ("STT", "idx"), ("Mã VT", "item"), ("Tên VT", "item_name"),
            ("SL", "qty"), ("UOM", "uom"), ("HĐ khung", "framework_contract"),
            ("Đơn giá ƯT (VND)", "estimated_unit_cost"), ("Thành tiền (VND)", "estimated_amount"),
            ("Kho đích", "warehouse"), ("Ngày cần", "schedule_date"), ("Ghi chú", "remarks"),
        ],
        "item_writable": {"item", "qty", "uom", "framework_contract", "estimated_unit_cost",
                          "warehouse", "schedule_date", "remarks"},
        "required_create": ["request_type", "transaction_date", "schedule_date"],
        "link_header": {"department": "SC Department", "warehouse": "SC Warehouse"},
        "link_item": {"item": "SC Item", "uom": "SC UOM", "framework_contract": "Framework Contract",
                      "warehouse": "SC Warehouse"},
        "editable_stage_field": None, "editable_stages": None,
    },
    "SC Purchase Order": {
        "label": "Đơn mua",
        "sheet_parent": "Phiếu",
        "header_cols": [
            ("Mã phiếu", "name"), ("NCC", "supplier"), ("Tên NCC", "supplier_name"),
            ("Ngày PO", "transaction_date"), ("Ngày giao DK", "schedule_date"), ("Kho nhận", "to_warehouse"),
            ("ĐK thanh toán", "payment_terms"), ("ĐK giao hàng", "delivery_terms"),
            ("HĐ khung", "framework_contract"), ("Tổng SL", "total_qty"), ("Tổng giá trị (VND)", "grand_total"),
            ("Trạng thái", "status"), ("Người tạo", "owner"), ("Ghi chú", "remarks"),
        ],
        "header_writable": {"supplier", "transaction_date", "schedule_date", "to_warehouse",
                            "payment_terms", "delivery_terms", "framework_contract", "remarks"},
        "child_field": "items", "child_doctype": "SC Purchase Order Item", "item_key": "item",
        "item_cols": [
            ("Mã phiếu", "parent"), ("STT", "idx"), ("Mã VT", "item"), ("Tên VT", "item_name"),
            ("SL", "qty"), ("UOM", "uom"), ("Đơn giá (VND)", "rate"), ("Thành tiền (VND)", "amount"),
            ("Kho", "warehouse"), ("Ngày cần", "schedule_date"),
        ],
        "item_writable": {"item", "qty", "uom", "rate", "warehouse", "schedule_date"},
        "required_create": ["supplier", "transaction_date", "schedule_date"],
        "link_header": {"supplier": "SC Supplier", "to_warehouse": "SC Warehouse",
                        "framework_contract": "Framework Contract"},
        "link_item": {"item": "SC Item", "uom": "SC UOM", "warehouse": "SC Warehouse"},
        "editable_stage_field": "approval_stage", "editable_stages": {"", "Draft", "Rejected"},
    },
    "SC Inventory Count Sheet": {
        "label": "Phiếu kiểm kê",
        "sheet_parent": "Phiếu",
        "header_cols": [
            ("Mã phiếu", "name"), ("Ngày kiểm kê", "count_date"), ("Kho", "warehouse"),
            ("Phạm vi", "count_scope"), ("Nhóm vật tư", "item_group"), ("Zone bin", "bin_zone"),
            ("Trạng thái", "status"), ("Tổng items", "total_items"), ("Người tạo", "owner"),
            ("Ghi chú", "remarks"),
        ],
        "header_writable": {"count_date", "warehouse", "count_scope", "item_group", "bin_zone", "remarks"},
        "child_field": "items", "child_doctype": "SC ICS Item", "item_key": "item",
        "item_cols": [
            ("Mã phiếu", "parent"), ("STT", "idx"), ("Mã VT", "item"), ("Tên VT", "item_name"),
            ("UOM", "uom"), ("Lô", "batch"), ("Bin", "bin_location"), ("SL hệ thống", "system_qty"),
            ("SL thực tế", "actual_qty"), ("Chênh lệch", "difference"), ("Lý do", "reason"),
            ("Ghi chú", "remarks"),
        ],
        "item_writable": {"item", "uom", "batch", "bin_location", "actual_qty", "reason", "remarks"},
        "required_create": ["count_date", "warehouse", "count_scope"],
        "require_items": False,  # ICS cho phép phiếu kiểm kê rỗng (snapshot tự điền)
        "link_header": {"warehouse": "SC Warehouse", "item_group": "SC Item Group"},
        "link_item": {"item": "SC Item", "uom": "SC UOM", "batch": "SC Batch", "bin_location": "Bin Location"},
        "editable_stage_field": None, "editable_stages": None,
    },
}


# Field SL chính của mỗi child (để CR-02 validate SL > 0).
ITEM_QTY_FIELD = {
    "Framework Contract": "contract_qty",
    "SC Material Request": "qty",
    "SC Purchase Order": "qty",
    "SC Purchase Receipt": "qty",
    "SC Transfer Request": "requested_qty",
    "SC Inventory Count Sheet": "actual_qty",
}
ALLOW_ZERO_QTY = {"SC Inventory Count Sheet"}  # kiểm kê: SL thực tế = 0 hợp lệ


def _cfg(doctype: str) -> dict:
    cfg = CONFIGS.get(doctype)
    if not cfg:
        frappe.throw(_("Doctype {0} chưa hỗ trợ xuất/nhập phiếu 2-sheet").format(doctype))
    return cfg


def _required_item_fields(doctype: str, cfg: dict) -> set[str]:
    """Field vật tư bắt buộc khi nhập — khớp đúng với parse_child_rows:
    Mã VT (item_key) + UOM (nếu writable) + SL chính (trừ doctype cho phép SL=0)."""
    req = {cfg["item_key"]}
    if "uom" in cfg["item_writable"]:
        req.add("uom")
    qtyf = ITEM_QTY_FIELD.get(doctype)
    if qtyf and doctype not in ALLOW_ZERO_QTY:
        req.add(qtyf)
    return req


def _mark_required(cols: list[tuple[str, str]], required: set[str]) -> list[str]:
    """Dựng dòng label, gắn ' *' vào cột bắt buộc. Chỉ áp dòng label (dòng 1);
    dòng fieldname (dòng 2) giữ nguyên → import round-trip an toàn."""
    return [(l + " *") if f in required else l for l, f in cols]


@frappe.whitelist()
def list_voucher_doctypes() -> list[str]:
    """Danh sách doctype hỗ trợ — frontend dùng để bật nút Xuất/Nhập 2-sheet."""
    return list(CONFIGS.keys())


# ---------------------------------------------------------------------------
# Helpers chung
# ---------------------------------------------------------------------------
def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        frappe.throw(_("Cần openpyxl để xử lý file Excel (.xlsx)"))


def _ftypes(doctype: str) -> dict[str, str]:
    return {df.fieldname: df.fieldtype for df in frappe.get_meta(doctype).fields}


def _fmt_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return value


def _coerce_date(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    s = str(value).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return str(getdate(s))


def _coerce_num(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(" ", "").replace(",", "")
    return flt(s)


def _coerce(value: Any, ftype: str) -> Any:
    """Coerce theo fieldtype thật. Ô rỗng → None (bỏ qua, không set)."""
    if value is None:
        return None
    if isinstance(value, str) and not value.strip():
        return None
    if ftype == "Date":
        return _coerce_date(value)
    if ftype == "Datetime":
        if isinstance(value, (datetime, date)):
            return str(value)
        return str(get_datetime(str(value).strip()))
    if ftype in ("Int", "Long Int"):
        return int(_coerce_num(value))
    if ftype in ("Float", "Currency", "Percent"):
        return _coerce_num(value)
    if ftype == "Check":
        if isinstance(value, (int, float)):
            return 1 if value else 0
        return 1 if str(value).strip().lower() in ("1", "true", "yes", "y", "x", "✓", "có", "co") else 0
    return value.strip() if isinstance(value, str) else str(value)


def _map_rows(rows2d: list[list], cols: list[tuple[str, str]]) -> list[dict]:
    """Map list-of-rows → list dict {fieldname: value}. Tự nhận diện header 1/2 dòng.
    Header chấp nhận cả fieldname lẫn label (lowercase)."""
    alias: dict[str, str] = {}
    fieldnames = {fn.lower() for _l, fn in cols}
    for label, fn in cols:
        alias[fn.lower()] = fn
        alias[label.strip().lower()] = fn

    def _norm(row):
        return [(str(c).strip() if c is not None else "") for c in row]

    rows = [_norm(r) for r in rows2d]
    if not rows:
        return []
    header_idx = 0
    if len(rows) >= 2:
        row2 = rows[1]
        nonempty = sum(1 for c in row2 if c)
        fn_hits = sum(1 for c in row2 if c.lower() in fieldnames)
        if nonempty and fn_hits >= max(2, (nonempty + 1) // 2):
            header_idx = 1
    headers = rows[header_idx]
    mapped = [alias.get(h.lower(), "") for h in headers]
    out: list[dict] = []
    for r in rows[header_idx + 1:]:
        if all(not c for c in r):
            continue
        rec: dict[str, Any] = {}
        for i, fn in enumerate(mapped):
            if fn:
                rec[fn] = r[i] if i < len(r) else None
        out.append(rec)
    return out


def _read_sheet(ws, cols: list[tuple[str, str]]) -> list[dict]:
    """Đọc 1 worksheet openpyxl → list dict (qua _map_rows)."""
    return _map_rows([list(r) for r in ws.iter_rows(values_only=True)], cols)


def _rows_from_file(content_b64: str, file_type: str) -> list[list]:
    """Đọc file CSV/XLSX (1 sheet đầu) → list-of-rows (cho CR-02 grid import)."""
    raw = base64.b64decode(content_b64)
    ft = (file_type or "xlsx").lower()
    if ft == "csv":
        import csv as _csv
        text = raw.decode("utf-8-sig", errors="replace")
        return [list(r) for r in _csv.reader(io.StringIO(text))]
    _require_openpyxl()
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    ws = wb.active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _guide_lines(cfg: dict) -> list[str]:
    key_label = cfg["header_cols"][0][0]  # vd "Mã HĐ" / "Mã phiếu"
    parent = cfg["sheet_parent"]
    return [
        f"HƯỚNG DẪN NHẬP: {cfg['label'].upper()}",
        "",
        f"File gồm 2 sheet: '{parent}' và '{SHEET_ITEM}', nối nhau bằng cột '{key_label}'.",
        "GIỮ NGUYÊN 2 dòng đầu mỗi sheet (dòng 1 = tên cột, dòng 2 = mã trường). Dữ liệu từ dòng 3.",
        "(Nếu bạn tự gõ file 1 dòng tiêu đề tiếng Việt cũng được — hệ thống tự nhận diện.)",
        "",
        f"► CỘT '{key_label}' — quan trọng:",
        f"   • TẠO MỚI: gõ tên bất kỳ (vd P1, A...). Hai sheet phải dùng CÙNG một chuỗi để",
        f"     khớp vật tư với phiếu. Hệ thống TỰ CẤP mã thật (KHÔNG dùng tên bạn gõ).",
        f"   • SỬA phiếu ĐÃ CÓ: dùng ĐÚNG mã của hệ thống. Cách chắc nhất: bấm 'Xuất' trước,",
        f"     sửa trên file đó rồi nhập lại.",
        f"   • Mã CHƯA có → tạo mới; ĐÃ có → cập nhật (thay toàn bộ danh mục vật tư).",
        "",
        "► Phiếu nhập vào LUÔN lưu ở dạng NHÁP (Draft) — chưa duyệt, chưa ghi sổ kho.",
        "► CHỈ phiếu đang Draft mới cập nhật/ghi đè được. Phiếu đã duyệt/đã ghi sổ sẽ bị BỎ QUA.",
        "",
        "► Cột chữ xám (Tên VT, Thành tiền, Tổng…, Người tạo, Trạng thái) chỉ để xem — không cần nhập.",
        "► Cột có dấu * ở tên cột là BẮT BUỘC khi tạo phiếu mới — không được bỏ trống.",
        "► Ngày dạng dd/mm/yyyy (vd 20/06/2026). Số tiền có thể có dấu phẩy (vd 2,848,000).",
    ]


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------
def _currency_cols(doctype: str, cols: list[tuple]) -> list[int]:
    """Chỉ số cột (1-based) là field Currency → áp định dạng tiền tệ VND."""
    cur = {f for f, t in _ftypes(doctype).items() if t == "Currency"}
    return [i for i, (_l, f) in enumerate(cols, start=1) if f in cur]


def _apply_vnd_format(ws, col_idxs: list[int], first_data_row: int):
    """Đặt number_format VND cho các cột tiền tệ, chỉ từ dòng dữ liệu trở đi."""
    for ci in col_idxs:
        for row in range(first_data_row, ws.max_row + 1):
            ws.cell(row=row, column=ci).number_format = VND_NUMFMT


def _build_workbook(doctype: str, cfg: dict, parents: list[dict],
                    items_by_parent: dict[str, list[dict]], guide: bool):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws_p = wb.active
    ws_p.title = cfg["sheet_parent"]
    ws_i = wb.create_sheet(SHEET_ITEM)
    bold = Font(bold=True)

    if guide:
        ws_g = wb.create_sheet(SHEET_GUIDE)
        for line in _guide_lines(cfg):
            ws_g.append([line])
        ws_g["A1"].font = Font(bold=True, size=13)
        ws_g.column_dimensions["A"].width = 95

    req_header = set(cfg.get("required_create", []))
    req_item = _required_item_fields(doctype, cfg)

    ws_p.append(_mark_required(cfg["header_cols"], req_header))
    ws_p.append([f for _l, f in cfg["header_cols"]])
    for c in ws_p[1]:
        c.font = bold
    for p in parents:
        ws_p.append([_fmt_cell(p.get(f)) for _l, f in cfg["header_cols"]])

    ws_i.append(_mark_required(cfg["item_cols"], req_item))
    ws_i.append([f for _l, f in cfg["item_cols"]])
    for c in ws_i[1]:
        c.font = bold
    item_count = 0
    for p in parents:
        for it in items_by_parent.get(p["name"], []):
            row = []
            for _l, f in cfg["item_cols"]:
                row.append(p["name"] if f == "parent" else _fmt_cell(it.get(f)))
            ws_i.append(row)
            item_count += 1

    # Định dạng tiền tệ VND cho các cột Currency (dữ liệu từ dòng 3; dòng 1-2 là header).
    _apply_vnd_format(ws_p, _currency_cols(doctype, cfg["header_cols"]), 3)
    _apply_vnd_format(ws_i, _currency_cols(cfg["child_doctype"], cfg["item_cols"]), 3)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read(), item_count


def _do_export(doctype: str, filters=None, order_by="modified desc",
               limit=10000, with_data: int = 1, guide: bool = False) -> dict:
    cfg = _cfg(doctype)
    _check_perm(doctype, "read")
    if isinstance(filters, str):
        filters = frappe.parse_json(filters) if filters.strip() else None

    parents: list[dict] = []
    items_by_parent: dict[str, list[dict]] = {}
    if cint(with_data):
        p_fields = [f for _l, f in cfg["header_cols"]]
        parents = frappe.get_all(doctype, fields=p_fields, filters=filters or {},
                                 order_by=order_by or "modified desc",
                                 limit=cint(limit) or 10000)
        names = [p["name"] for p in parents]
        if names:
            i_fields = ["parent"] + [f for _l, f in cfg["item_cols"] if f != "parent"]
            rows = frappe.get_all(cfg["child_doctype"], fields=i_fields,
                                  filters={"parent": ["in", names], "parenttype": doctype},
                                  order_by="parent asc, idx asc", limit=0)
            for r in rows:
                items_by_parent.setdefault(r["parent"], []).append(r)

    content, item_count = _build_workbook(doctype, cfg, parents, items_by_parent, guide=guide)
    suffix = "" if cint(with_data) else "_template"
    return {
        "filename": f"{_safe_filename(doctype)}{suffix}.xlsx",
        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "content_b64": base64.b64encode(content).decode("ascii"),
        "contract_count": len(parents),
        "item_count": item_count,
    }


@frappe.whitelist()
def export_voucher(doctype: str, filters=None, order_by: str = "modified desc",
                   limit: int | str = 10000) -> dict:
    """Xuất danh sách phiếu + danh mục vật tư ra Excel 2 sheet."""
    _require_openpyxl()
    return _do_export(doctype, filters=filters, order_by=order_by, limit=limit, with_data=1)


@frappe.whitelist()
def voucher_template(doctype: str, with_data: int | str = 0, limit: int | str = 50) -> dict:
    """Tải template Excel 2 sheet (kèm sheet Hướng dẫn)."""
    _require_openpyxl()
    return _do_export(doctype, filters=None, order_by="modified desc",
                      limit=limit, with_data=cint(with_data), guide=True)


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------
def _build_payload(cfg: dict, ft_p: dict, ft_c: dict, header: dict | None,
                   items: list[dict]) -> dict:
    payload: dict[str, Any] = {}
    header = header or {}
    for fn in cfg["header_writable"]:
        if fn in header:
            v = _coerce(header.get(fn), ft_p.get(fn, "Data"))
            if v is not None:
                payload[fn] = v
    keyf = cfg["item_key"]
    child_rows: list[dict] = []
    for it in items:
        code = it.get(keyf)
        code = str(code).strip() if code not in (None, "") else None
        if not code:
            continue
        row: dict[str, Any] = {}
        for fn in cfg["item_writable"]:
            if fn in it:
                v = _coerce(it.get(fn), ft_c.get(fn, "Data"))
                if v is not None:
                    row[fn] = v
        row[keyf] = code
        child_rows.append(row)
    payload["__items__"] = child_rows
    return payload


def _validate_links(cfg: dict, payload: dict, errors: list[str], key: str):
    lbl = cfg["label"]
    for fn, dt in cfg["link_header"].items():
        v = payload.get(fn)
        if v and not frappe.db.exists(dt, v):
            errors.append(_("{0} {1}: {2} '{3}' không tồn tại").format(lbl, key, fn, v))
    for it in payload["__items__"]:
        for fn, dt in cfg["link_item"].items():
            v = it.get(fn)
            if v and not frappe.db.exists(dt, v):
                errors.append(_("{0} {1}: {2} '{3}' không tồn tại trong {4}").format(
                    lbl, key, fn, v, dt))


def _editable(cfg: dict, db: dict) -> tuple[bool, str]:
    if cint(db.get("docstatus")) != 0:
        return False, _("đã submit/ghi sổ (docstatus={0})").format(db.get("docstatus"))
    sf = cfg.get("editable_stage_field")
    if sf and (db.get(sf) or "") not in cfg["editable_stages"]:
        return False, _("đang ở stage {0}").format(db.get(sf))
    return True, ""


@frappe.whitelist()
def import_voucher(doctype: str, content_b64: str, dry_run: int | str = 1,
                   allow_create: int | str = 1) -> dict:
    """Nhập phiếu + danh mục vật tư từ Excel 2 sheet. Phiếu nhập LUÔN ở Draft.

    - Gom theo Mã phiếu. Đã có → cập nhật (chỉ khi Draft); chưa có → tạo mới Draft.
    - dry_run=1: chỉ validate + đếm; =0: commit (savepoint mỗi phiếu).
    """
    cfg = _cfg(doctype)
    _require_openpyxl()
    _check_perm(doctype, "create")
    _check_perm(doctype, "write")
    is_dry = bool(cint(dry_run))
    can_create = bool(cint(allow_create))
    ft_p = _ftypes(doctype)
    ft_c = _ftypes(cfg["child_doctype"])

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(base64.b64decode(content_b64)), data_only=True, read_only=True)
    titles = {t.lower(): t for t in wb.sheetnames}
    ws_p = wb[titles[cfg["sheet_parent"].lower()]] if cfg["sheet_parent"].lower() in titles else None
    ws_i = wb[titles[SHEET_ITEM.lower()]] if SHEET_ITEM.lower() in titles else None
    if ws_p is None and ws_i is None:
        frappe.throw(_("File phải có sheet '{0}' và/hoặc '{1}'").format(cfg["sheet_parent"], SHEET_ITEM))

    parent_rows = _read_sheet(ws_p, cfg["header_cols"]) if ws_p is not None else []
    item_rows = _read_sheet(ws_i, cfg["item_cols"]) if ws_i is not None else []

    parents_by_key: dict[str, dict] = {}
    items_by_key: dict[str, list[dict]] = {}
    order: list[str] = []
    for r in parent_rows:
        k = r.get("name")
        k = str(k).strip() if k not in (None, "") else None
        if not k:
            continue
        if k not in parents_by_key:
            order.append(k)
        parents_by_key[k] = r
    for r in item_rows:
        k = r.get("parent")
        k = str(k).strip() if k not in (None, "") else None
        if not k:
            continue
        if k not in parents_by_key and k not in items_by_key:
            order.append(k)
        items_by_key.setdefault(k, []).append(r)

    summary = {"total": len(order), "created": 0, "updated": 0, "skipped": 0,
               "failed": 0, "would_create": 0, "would_update": 0}
    preview: list[dict] = []
    errors: list[str] = []
    stage_field = cfg.get("editable_stage_field")

    for key in order:
        header = parents_by_key.get(key)
        items = items_by_key.get(key, [])
        exists = bool(frappe.db.exists(doctype, key))

        err_before = len(errors)
        payload = _build_payload(cfg, ft_p, ft_c, header, items)
        _validate_links(cfg, payload, errors, key)
        has_err = len(errors) > err_before

        if exists:
            cols = ["docstatus"] + ([stage_field] if stage_field else [])
            db = frappe.db.get_value(doctype, key, cols, as_dict=True) or {}
            ok, reason = _editable(cfg, db)
            if not ok:
                summary["skipped"] += 1
                preview.append({"key": key, "name": key, "action": "skipped",
                                "item_count": len(payload["__items__"]),
                                "reason": _("Phiếu không sửa được — {0}").format(reason)})
                continue
            action = "update"
        else:
            if not can_create:
                summary["skipped"] += 1
                preview.append({"key": key, "name": key, "action": "skipped",
                                "item_count": len(payload["__items__"]),
                                "reason": _("Phiếu chưa tồn tại (không bật tạo mới)")})
                continue
            missing = [fn for fn in cfg["required_create"] if not payload.get(fn)]
            if missing:
                errors.append(_("{0} {1}: tạo mới thiếu {2}").format(cfg["label"], key, ", ".join(missing)))
                has_err = True
            if cfg.get("require_items", True) and not payload["__items__"]:
                errors.append(_("{0} {1}: tạo mới cần ≥1 dòng vật tư").format(cfg["label"], key))
                has_err = True
            action = "create"

        if has_err:
            summary["failed"] += 1
            preview.append({"key": key, "name": key, "action": "error",
                            "item_count": len(payload["__items__"]),
                            "error": _("Có lỗi dữ liệu (xem danh sách lỗi)")})
            continue

        if is_dry:
            summary["would_create" if action == "create" else "would_update"] += 1
            preview.append({"key": key, "name": key if exists else _("(mã mới)"),
                            "action": action, "item_count": len(payload["__items__"])})
            continue

        sp = f"voucher_imp_{abs(hash(key)) % 10_000_000}"
        try:
            frappe.db.savepoint(sp)
            if action == "update":
                doc = frappe.get_doc(doctype, key)
            else:
                doc = frappe.new_doc(doctype)
            for fn, val in payload.items():
                if fn == "__items__":
                    continue
                doc.set(fn, val)
            doc.set(cfg["child_field"], payload["__items__"])  # thay toàn bộ danh mục
            if action == "update":
                doc.save()  # GIỮ Draft — không submit
            else:
                doc.insert()  # tạo mới Draft
            summary["updated" if action == "update" else "created"] += 1
            preview.append({"key": key, "name": doc.name,
                            "action": "updated" if action == "update" else "created",
                            "item_count": len(doc.get(cfg["child_field"]))})
        except Exception as e:
            try:
                frappe.db.rollback(save_point=sp)
            except Exception:
                pass
            summary["failed"] += 1
            msg = str(e)[:400]
            errors.append(_("{0} {1}: {2}").format(cfg["label"], key, msg))
            preview.append({"key": key, "name": key, "action": "failed",
                            "item_count": len(payload["__items__"]), "error": msg})

    if not is_dry:
        frappe.db.commit()

    return {"dry_run": is_dry, "summary": summary, "preview": preview, "errors": errors}


# ---------------------------------------------------------------------------
# CR-02 — Import/Export/Template ngay tại lưới "Danh mục vật tư" trong form
#   (parse + validate, KHÔNG ghi DB — frontend nạp vào lưới in-memory)
# ---------------------------------------------------------------------------
def _child_grid_cols(cfg: dict) -> list[tuple[str, str]]:
    """Cột cho lưới (bỏ parent + idx — không cần khi nhập trong form)."""
    return [(l, f) for l, f in cfg["item_cols"] if f not in ("parent", "idx")]


@frappe.whitelist()
def child_template(doctype: str, file_type: str = "xlsx") -> dict:
    """Tải file mẫu CHỈ cột vật tư (child) cho lưới trong form — CR-02."""
    cfg = _cfg(doctype)
    _check_perm(doctype, "read")
    cols = _child_grid_cols(cfg)
    labels = _mark_required(cols, _required_item_fields(doctype, cfg))
    fields = [f for _l, f in cols]
    fname = f"{_safe_filename(cfg['child_doctype'])}_template"

    if (file_type or "xlsx").lower() == "csv":
        import csv as _csv
        buf = io.StringIO()
        w = _csv.writer(buf)
        w.writerow(labels)
        w.writerow(fields)
        text = "﻿" + buf.getvalue()
        return {"filename": f"{fname}.csv", "content_type": "text/csv; charset=utf-8",
                "content_b64": base64.b64encode(text.encode("utf-8")).decode("ascii")}

    _require_openpyxl()
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); ws = wb.active; ws.title = SHEET_ITEM
    ws.append(labels); ws.append(fields)
    for c in ws[1]:
        c.font = Font(bold=True)
    b = io.BytesIO(); wb.save(b); b.seek(0)
    return {"filename": f"{fname}.xlsx",
            "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "content_b64": base64.b64encode(b.read()).decode("ascii")}


@frappe.whitelist()
def parse_child_rows(doctype: str, content_b64: str, file_type: str = "xlsx") -> dict:
    """Parse + validate file vật tư cho lưới trong form (CR-02). KHÔNG ghi DB.

    Trả {total, ok_count, error_count, rows_ok, rows_error}:
      - rows_ok: list dict field con (đã coerce) + item_name → frontend nạp vào lưới.
      - rows_error: [{line, item, errors[]}] để hiển thị, KHÔNG nạp dòng lỗi.
    """
    cfg = _cfg(doctype)
    _check_perm(doctype, "read")
    ft_c = _ftypes(cfg["child_doctype"])
    keyf = cfg["item_key"]
    qtyf = ITEM_QTY_FIELD.get(doctype)
    allow_zero = doctype in ALLOW_ZERO_QTY

    recs = _map_rows(_rows_from_file(content_b64, file_type), cfg["item_cols"])
    rows_ok: list[dict] = []
    rows_error: list[dict] = []
    for i, rec in enumerate(recs, start=1):
        errs: list[str] = []
        code = rec.get(keyf)
        code = str(code).strip() if code not in (None, "") else None

        out: dict[str, Any] = {}
        for fn in cfg["item_writable"]:
            if fn in rec:
                v = _coerce(rec.get(fn), ft_c.get(fn, "Data"))
                if v is not None:
                    out[fn] = v
        if code:
            out[keyf] = code
        else:
            errs.append(_("thiếu Mã VT"))

        for fn, dt in cfg["link_item"].items():
            v = out.get(fn)
            if v and not frappe.db.exists(dt, v):
                errs.append(_("{0} '{1}' không tồn tại").format(fn, v))

        if "uom" in cfg["item_writable"] and not out.get("uom"):
            errs.append(_("thiếu UOM"))
        if qtyf and not allow_zero and flt(out.get(qtyf) or 0) <= 0:
            errs.append(_("SL phải > 0"))

        if errs:
            rows_error.append({"line": i, "item": code or "", "errors": errs})
        else:
            nm = frappe.db.get_value("SC Item", code, "item_name")
            if nm:
                out["item_name"] = nm
            rows_ok.append(out)

    return {"total": len(recs), "ok_count": len(rows_ok),
            "error_count": len(rows_error), "rows_ok": rows_ok, "rows_error": rows_error}
