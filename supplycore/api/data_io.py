"""SupplyCore Data Import/Export — Frappe-style import/export trên SPA.

Hoạt động giống Frappe Data Import:
  - Tải template CSV/XLSX (header = fieldname, hoặc có sẵn dữ liệu để chỉnh sửa).
  - Export dữ liệu hiện tại ra CSV/XLSX (kèm filter).
  - Upload → dry-run (preview + báo lỗi) → commit (insert/update/skip).
  - Hỗ trợ child table (Table fieldtype) thông qua cột JSON-array.
  - Tôn trọng role permission (create/write/submit).
  - Partial-success: row lỗi rollback bằng savepoint, row OK vẫn commit.

API path: supplycore.api.data_io.*
"""

from __future__ import annotations

import base64
import csv
import io
import json
from typing import Any

import frappe
from frappe import _
from frappe.utils import cint, flt, getdate, get_datetime

# Fieldtype không xuất/không nhập
SKIP_FIELDTYPES = {
    "Section Break", "Column Break", "Tab Break", "HTML", "Button",
    "Heading", "Fold", "Image",
}
# Field hệ thống bỏ qua khi import
SYSTEM_FIELDS = {
    "name", "owner", "creation", "modified", "modified_by", "docstatus",
    "idx", "parent", "parentfield", "parenttype", "lft", "rgt",
    "old_parent", "_user_tags", "_assign", "_liked_by", "_comments",
}


# Định dạng tiền tệ VND cho ô Excel: "1.234.567 ₫", 0 chữ số thập phân.
# Giá trị ô vẫn là SỐ → đọc lại (data_only) không vỡ round-trip import.
VND_NUMFMT = '#,##0" ₫"'


def _currency_fields(doctype: str) -> set[str]:
    return {df.fieldname for df in frappe.get_meta(doctype).fields
            if df.fieldtype == "Currency"}


def _currency_idx(doctype: str, headers: list[str]) -> set[int]:
    """Chỉ số cột (0-based) là field Currency."""
    cur = _currency_fields(doctype)
    return {i for i, h in enumerate(headers) if h in cur}


def _vnd_int(value: Any) -> str:
    """Chuỗi số nguyên VND cho CSV (không phân tách nghìn → nhập lại an toàn)."""
    if value in (None, ""):
        return ""
    return str(int(round(flt(value))))


def _check_perm(doctype: str, perm: str) -> None:
    if not frappe.has_permission(doctype, perm):
        frappe.throw(_("Không có quyền {0} {1}").format(perm, doctype),
                     frappe.PermissionError)


def _serialize_df(df) -> dict:
    return {
        "fieldname": df.fieldname,
        "label": df.label or df.fieldname,
        "fieldtype": df.fieldtype,
        "options": df.options,
        "reqd": int(df.reqd or 0),
        "unique": int(df.unique or 0),
        "readonly": int(df.read_only or 0),
        "default": df.default,
        "in_list_view": int(df.in_list_view or 0),
    }


def _meta_fields(doctype: str, include_child: bool = True) -> list[dict]:
    meta = frappe.get_meta(doctype)
    out: list[dict] = []
    for df in meta.fields:
        if df.fieldtype in SKIP_FIELDTYPES:
            continue
        if df.fieldtype == "Table":
            child = _serialize_df(df)
            if include_child:
                child["child_fields"] = _meta_fields(df.options, include_child=False)
            out.append(child)
        else:
            out.append(_serialize_df(df))
    return out


def _writable_field_names(doctype: str, include_table: bool = True) -> list[str]:
    """Field có thể set qua import (loại read-only + system fields)."""
    meta = frappe.get_meta(doctype)
    names: list[str] = []
    for df in meta.fields:
        if df.fieldtype in SKIP_FIELDTYPES:
            continue
        if df.fieldname in SYSTEM_FIELDS:
            continue
        if df.fieldtype == "Table":
            if include_table:
                names.append(df.fieldname)
            continue
        if df.read_only:
            continue
        names.append(df.fieldname)
    return names


def _exportable_field_names(doctype: str) -> list[str]:
    """Field có thể export — include name + readonly fields cho audit."""
    meta = frappe.get_meta(doctype)
    names: list[str] = ["name"]
    for df in meta.fields:
        if df.fieldtype in SKIP_FIELDTYPES:
            continue
        if df.fieldname in SYSTEM_FIELDS:
            continue
        if df.fieldtype == "Table":
            continue  # child tables exported as JSON in template
        names.append(df.fieldname)
    return names


def _fieldname_map(doctype: str) -> dict[str, str]:
    """Map cả fieldname VÀ label (lowercase) → fieldname. Cho phép header
    template dùng tiếng Việt (label) hoặc fieldname."""
    meta = frappe.get_meta(doctype)
    m: dict[str, str] = {"name": "name"}
    for df in meta.fields:
        if df.fieldtype in SKIP_FIELDTYPES:
            continue
        m[df.fieldname.lower()] = df.fieldname
        if df.label:
            m[df.label.strip().lower()] = df.fieldname
    return m


# ---------------------------------------------------------------------------
# Schema + listing
# ---------------------------------------------------------------------------

@frappe.whitelist()
def get_doctype_schema(doctype: str) -> dict:
    """Trả meta fields + flags cho UI preview."""
    _check_perm(doctype, "read")
    meta = frappe.get_meta(doctype)
    return {
        "doctype": doctype,
        "autoname": meta.autoname,
        "title_field": meta.title_field,
        "is_submittable": int(meta.is_submittable or 0),
        "is_tree": int(getattr(meta, "is_tree", 0) or 0),
        "fields": _meta_fields(doctype),
        "writable_fields": _writable_field_names(doctype),
        "exportable_fields": _exportable_field_names(doctype),
    }


# ---------------------------------------------------------------------------
# Template + export
# ---------------------------------------------------------------------------

def _row_to_strings(row: dict, headers: list[str],
                    currency_set: set[str] | None = None) -> list[Any]:
    """Chuyển record → list ô. Ô tiền tệ giữ dạng SỐ (để XLSX format VND +
    round-trip), ô khác → chuỗi. currency_set=None → mọi ô là chuỗi (cũ)."""
    currency_set = currency_set or set()
    out: list[Any] = []
    for h in headers:
        v = row.get(h)
        if h in currency_set:
            out.append(None if v in (None, "") else flt(v))
        else:
            out.append("" if v is None else str(v))
    return out


def _build_file(filename_base: str, headers: list[str], rows: list[list[Any]],
                file_type: str, label_row: list[str] | None = None,
                currency_idx: set[int] | None = None) -> dict:
    """Build CSV/XLSX file.

    label_row (optional): nếu truyền, sẽ ghi LÀM DÒNG 1 (tên hiển thị frontend),
    headers thành dòng 2 (fieldname), data từ dòng 3. Đây là format 2-dòng-header
    dùng cho list import/export inline.

    currency_idx (optional): chỉ số cột (0-based) là tiền tệ. XLSX → number_format
    VND (ô vẫn là số); CSV → ghi số nguyên VND (round-trip an toàn).
    """
    currency_idx = currency_idx or set()
    file_type = (file_type or "csv").lower()
    if file_type == "xlsx":
        try:
            from openpyxl import Workbook
        except ImportError:
            frappe.throw(_("Cần openpyxl để xuất XLSX"))
        wb = Workbook()
        ws = wb.active
        ws.title = filename_base[:31] or "Sheet1"
        if label_row is not None:
            ws.append(label_row)
        ws.append(headers)
        for r in rows:
            ws.append(r)
        if currency_idx:
            first_data = 3 if label_row is not None else 2
            for ci in currency_idx:
                for rn in range(first_data, ws.max_row + 1):
                    ws.cell(row=rn, column=ci + 1).number_format = VND_NUMFMT
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return {
            "filename": f"{filename_base}.xlsx",
            "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "content_b64": base64.b64encode(buf.read()).decode("ascii"),
        }
    # CSV (default) — UTF-8 BOM để Excel đọc tiếng Việt đúng
    def _csv_row(r: list[Any]) -> list[Any]:
        if not currency_idx:
            return r
        return [_vnd_int(v) if i in currency_idx else v for i, v in enumerate(r)]

    buf = io.StringIO()
    w = csv.writer(buf, quoting=csv.QUOTE_MINIMAL)
    if label_row is not None:
        w.writerow(label_row)
    w.writerow(headers)
    for r in rows:
        w.writerow(_csv_row(r))
    text = "﻿" + buf.getvalue()
    return {
        "filename": f"{filename_base}.csv",
        "content_type": "text/csv; charset=utf-8",
        "content_b64": base64.b64encode(text.encode("utf-8")).decode("ascii"),
    }


def _label_for_fields(doctype: str, fields: list[str]) -> list[str]:
    """Trả label hiển thị frontend cho mỗi fieldname (dòng 1 của file).

    Trường bắt buộc (reqd) được gắn dấu ' *' cuối label để người nhập biết
    cột nào không được bỏ trống. Dấu chỉ nằm ở dòng label (dòng 1) — dòng
    fieldname (dòng 2) giữ nguyên nên import round-trip không bị ảnh hưởng.
    """
    meta = frappe.get_meta(doctype)
    label_map = {df.fieldname: (df.label or df.fieldname) for df in meta.fields}
    reqd_set = {df.fieldname for df in meta.fields if df.reqd}  # 'name' không tính (autoname)
    label_map.setdefault("name", _("Mã / ID"))
    return [(label_map.get(f, f) + " *") if f in reqd_set else label_map.get(f, f)
            for f in fields]


def _safe_filename(doctype: str) -> str:
    return doctype.replace(" ", "_").replace("/", "_")


@frappe.whitelist()
def get_template(doctype: str, file_type: str = "csv",
                  with_data: int | str = 0, limit: int | str = 100) -> dict:
    """Tải template import. with_data=1 → kèm dữ liệu hiện tại để chỉnh sửa.

    Header = fieldname (round-trip an toàn). Child table xuất dưới dạng JSON array.
    """
    _check_perm(doctype, "read")
    meta = frappe.get_meta(doctype)

    flat_fields = _writable_field_names(doctype, include_table=False)
    child_table_fields = [df for df in meta.fields if df.fieldtype == "Table"
                           and df.fieldname not in SYSTEM_FIELDS]
    child_names = [df.fieldname for df in child_table_fields]
    headers = ["name"] + flat_fields + child_names

    rows: list[list[Any]] = []
    if cint(with_data):
        records = frappe.get_all(
            doctype,
            fields=["name"] + flat_fields,
            limit=cint(limit) or 100,
            order_by="modified desc",
        )
        # Pre-fetch child rows cho từng record (1 query / child table)
        child_data_by_record: dict[str, dict[str, list]] = {}
        if records and child_table_fields:
            names = [r["name"] for r in records]
            for cdf in child_table_fields:
                child_doctype = cdf.options
                child_meta = frappe.get_meta(child_doctype)
                child_field_list = [c.fieldname for c in child_meta.fields
                                     if c.fieldtype not in SKIP_FIELDTYPES
                                     and c.fieldname not in SYSTEM_FIELDS
                                     and c.fieldtype != "Table"]
                child_rows = frappe.get_all(
                    child_doctype,
                    filters={"parent": ["in", names],
                              "parenttype": doctype,
                              "parentfield": cdf.fieldname},
                    fields=["parent", "idx"] + child_field_list,
                    order_by="parent asc, idx asc",
                )
                for cr in child_rows:
                    pname = cr.pop("parent")
                    cr.pop("idx", None)
                    child_data_by_record.setdefault(pname, {}).setdefault(cdf.fieldname, []).append(cr)

        cur_set = _currency_fields(doctype)
        for r in records:
            row: list[Any] = []
            for h in headers:
                if h in child_names:
                    items = child_data_by_record.get(r["name"], {}).get(h, [])
                    row.append(json.dumps(items, default=str, ensure_ascii=False) if items else "")
                elif h in cur_set:
                    v = r.get(h)
                    row.append(None if v in (None, "") else flt(v))
                else:
                    v = r.get(h)
                    row.append("" if v is None else str(v))
            rows.append(row)

    return _build_file(_safe_filename(doctype), headers, rows, file_type,
                       currency_idx=_currency_idx(doctype, headers))


@frappe.whitelist()
def export_data(doctype: str, fields=None, filters=None,
                file_type: str = "csv", limit: int | str = 10000,
                order_by: str = "modified desc") -> dict:
    """Export rows ra CSV/XLSX. Tuỳ chọn fields + filters (1-dòng header)."""
    _check_perm(doctype, "read")
    if isinstance(fields, str):
        fields = json.loads(fields) if fields.strip() else None
    if isinstance(filters, str):
        filters = json.loads(filters) if filters.strip() else None
    fields = fields or _exportable_field_names(doctype)
    # Always include name for round-trip
    if "name" not in fields:
        fields = ["name"] + list(fields)

    records = frappe.get_all(doctype, fields=fields, filters=filters or {},
                              order_by=order_by, limit=cint(limit) or 10000)
    cur_set = _currency_fields(doctype)
    rows = [_row_to_strings(r, fields, cur_set) for r in records]
    return _build_file(_safe_filename(doctype), fields, rows, file_type,
                       currency_idx=_currency_idx(doctype, fields))


@frappe.whitelist()
def export_list(doctype: str, columns=None, filters=None,
                file_type: str = "csv", limit: int | str = 10000,
                order_by: str = "modified desc") -> dict:
    """Export theo list view với format 3-dòng (inline import/export):
      - Dòng 1 = tên hiển thị frontend (label)
      - Dòng 2 = fieldname (name — dùng để round-trip import)
      - Dòng 3+ = dữ liệu

    Args:
      columns: JSON list fieldname user chọn. None → tất cả exportable field.
      filters: Frappe filters (dict hoặc list) — lấy từ filter list đang hiển thị.
      order_by: sort hiện tại của list.
    """
    _check_perm(doctype, "read")
    if isinstance(columns, str):
        columns = json.loads(columns) if columns.strip() else None
    if isinstance(filters, str):
        filters = json.loads(filters) if filters.strip() else None

    columns = columns or _exportable_field_names(doctype)
    # Luôn có 'name' đầu tiên để import round-trip
    if "name" not in columns:
        columns = ["name"] + list(columns)
    else:
        # đảm bảo name ở đầu
        columns = ["name"] + [c for c in columns if c != "name"]

    # Loại field không tồn tại (an toàn)
    meta = frappe.get_meta(doctype)
    valid = {df.fieldname for df in meta.fields} | {"name"}
    columns = [c for c in columns if c in valid]

    label_row = _label_for_fields(doctype, columns)

    records = frappe.get_all(doctype, fields=columns, filters=filters or {},
                              order_by=order_by or "modified desc",
                              limit=cint(limit) or 10000)
    cur_set = _currency_fields(doctype)
    rows = [_row_to_strings(r, columns, cur_set) for r in records]
    out = _build_file(_safe_filename(doctype), columns, rows, file_type,
                       label_row=label_row,
                       currency_idx=_currency_idx(doctype, columns))
    out["row_count"] = len(rows)
    out["columns"] = columns
    return out


@frappe.whitelist()
def get_list_columns(doctype: str) -> dict:
    """Trả danh sách cột chọn được cho export — {fieldname, label, in_list_view}.
    UI dùng để render checkbox chọn cột."""
    _check_perm(doctype, "read")
    meta = frappe.get_meta(doctype)
    cols = [{"fieldname": "name", "label": _("Mã / ID"), "in_list_view": 1}]
    for df in meta.fields:
        if df.fieldtype in SKIP_FIELDTYPES or df.fieldtype == "Table":
            continue
        if df.fieldname in SYSTEM_FIELDS:
            continue
        cols.append({
            "fieldname": df.fieldname,
            "label": df.label or df.fieldname,
            "in_list_view": int(df.in_list_view or 0),
            "fieldtype": df.fieldtype,
        })
    return {"doctype": doctype, "columns": cols}


@frappe.whitelist()
def get_list_template(doctype: str, columns=None, file_type: str = "csv",
                       with_data: int | str = 0, limit: int | str = 50) -> dict:
    """Tải template import format 3-dòng:
      - Dòng 1 = tên hiển thị (label)
      - Dòng 2 = fieldname
      - Dòng 3+ = rỗng (template trống) hoặc dữ liệu hiện có (với with_data=1)

    columns: JSON list fieldname. None → tất cả field WRITABLE (có thể nhập).
             Luôn kèm 'name' đầu tiên để update round-trip.
    """
    _check_perm(doctype, "read")
    if isinstance(columns, str):
        columns = json.loads(columns) if columns.strip() else None

    if not columns:
        # Mặc định: name + tất cả field writable (loại Table)
        columns = ["name"] + _writable_field_names(doctype, include_table=False)
    else:
        columns = ["name"] + [c for c in columns if c != "name"]

    meta = frappe.get_meta(doctype)
    valid = {df.fieldname for df in meta.fields} | {"name"}
    columns = [c for c in columns if c in valid]

    label_row = _label_for_fields(doctype, columns)

    rows: list[list[Any]] = []
    if cint(with_data):
        records = frappe.get_all(doctype, fields=columns, filters={},
                                  order_by="modified desc",
                                  limit=cint(limit) or 50)
        rows = [_row_to_strings(r, columns) for r in records]

    out = _build_file(_safe_filename(doctype) + "_template", columns, rows,
                       file_type, label_row=label_row)
    out["row_count"] = len(rows)
    out["columns"] = columns
    return out


# ---------------------------------------------------------------------------
# Parse upload
# ---------------------------------------------------------------------------

def _parse_payload(content_b64: str, file_type: str,
                   two_row: bool = False) -> tuple[list[str], list[dict], int]:
    """Parse CSV/XLSX → (headers, data_rows, data_start_line).

    two_row=False: dòng 1 = header, data từ dòng 2 → data_start_line=2.
    two_row=True : dòng 1 = label (BỎ QUA), dòng 2 = fieldname header,
                   data từ dòng 3 → data_start_line=3.
    """
    raw = base64.b64decode(content_b64)
    file_type = (file_type or "csv").lower()

    def _empty(r):
        return all(v is None or (isinstance(v, str) and not v.strip()) for v in r)

    if file_type == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            frappe.throw(_("Cần openpyxl để đọc XLSX"))
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            if two_row:
                next(rows_iter)            # dòng 1 = label → bỏ qua
            header_row = next(rows_iter)   # dòng header (fieldname)
        except StopIteration:
            return [], [], (3 if two_row else 2)
        headers = [(str(h).strip() if h is not None else "") for h in header_row]
        data: list[dict] = []
        for r in rows_iter:
            if _empty(r):
                continue
            data.append({headers[i]: r[i] for i in range(min(len(headers), len(r)))})
        return headers, data, (3 if two_row else 2)

    # CSV
    text = raw.decode("utf-8-sig", errors="replace")
    all_rows = list(csv.reader(io.StringIO(text)))
    if not all_rows:
        return [], [], (3 if two_row else 2)
    header_idx = 1 if two_row else 0
    if len(all_rows) <= header_idx:
        return [], [], (3 if two_row else 2)
    headers = [(h.strip() if h else "") for h in all_rows[header_idx]]
    data = []
    for r in all_rows[header_idx + 1:]:
        if _empty(r):
            continue
        data.append({headers[i]: (r[i] if i < len(r) else None)
                     for i in range(len(headers))})
    return headers, data, (3 if two_row else 2)


def _coerce_value(value: Any, df) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        s = value.strip()
        if s == "":
            return None
    else:
        s = value
    ft = df.fieldtype
    if ft in ("Int", "Long Int"):
        return int(float(s)) if not isinstance(s, int) else s
    if ft in ("Float", "Currency", "Percent"):
        return flt(s)
    if ft == "Check":
        if isinstance(s, (int, float)):
            return 1 if s else 0
        return 1 if str(s).strip().lower() in ("1", "true", "yes", "y", "x", "✓", "có") else 0
    if ft == "Date":
        return str(getdate(s))
    if ft == "Datetime":
        return str(get_datetime(s))
    return s if isinstance(s, str) else str(s)


def _build_doc_dict(doctype: str, row: dict, fieldname_map: dict[str, str],
                     errors: list[str], line_no: int) -> dict:
    """Parse 1 row → doc dict. Validate Link existence + coerce types."""
    meta = frappe.get_meta(doctype)
    field_map = {df.fieldname: df for df in meta.fields}
    doc: dict[str, Any] = {"doctype": doctype}

    for hdr, val in row.items():
        if hdr is None:
            continue
        key = str(hdr).strip().lower()
        if not key:
            continue
        fn = fieldname_map.get(key)
        if not fn:
            continue
        if fn == "name":
            if val not in (None, "") and not (isinstance(val, str) and not val.strip()):
                doc["name"] = str(val).strip()
            continue
        if fn not in field_map:
            continue
        df = field_map[fn]
        if df.fieldtype == "Table":
            if val in (None, ""):
                continue
            try:
                items = json.loads(val) if isinstance(val, str) else val
                if isinstance(items, list):
                    doc[fn] = items
                else:
                    errors.append(_("Dòng {0}: '{1}' phải là JSON array").format(line_no, fn))
            except Exception as e:
                errors.append(_("Dòng {0}: {1} JSON không hợp lệ — {2}").format(line_no, fn, e))
            continue
        try:
            doc[fn] = _coerce_value(val, df)
        except Exception as e:
            errors.append(_("Dòng {0}: {1} = '{2}' không hợp lệ ({3})").format(
                line_no, df.label or fn, val, e))
            continue
        # Link existence
        if df.fieldtype == "Link" and doc.get(fn):
            if not frappe.db.exists(df.options, doc[fn]):
                errors.append(_("Dòng {0}: {1} = '{2}' không tồn tại trong {3}").format(
                    line_no, df.label or fn, doc[fn], df.options))
    return doc


# ---------------------------------------------------------------------------
# Dry-run + commit
# ---------------------------------------------------------------------------

@frappe.whitelist()
def import_data(doctype: str, content_b64: str, file_type: str = "csv",
                update_existing: int | str = 0, submit_after: int | str = 0,
                dry_run: int | str = 1, two_row_header: int | str = 0) -> dict:
    """Import data từ CSV/XLSX (base64). dry_run=1 → chỉ validate; =0 → commit.

    two_row_header=1: file có 2 dòng header (dòng 1 = tên hiển thị, dòng 2 =
    fieldname); dữ liệu bắt đầu dòng 3. Dùng cho list import/export inline.
    """
    # Chặn import vào sổ kho / sổ kế toán — chỉ được sinh qua voucher.
    if doctype in EXPORT_ONLY:
        frappe.throw(_("{0} là sổ hệ thống — chỉ cho phép EXPORT, không nhập trực tiếp"
                        ).format(doctype))
    _check_perm(doctype, "create")
    if cint(update_existing):
        _check_perm(doctype, "write")
    if cint(submit_after):
        _check_perm(doctype, "submit")

    two_row = bool(cint(two_row_header))
    headers, rows, data_start = _parse_payload(content_b64, file_type, two_row=two_row)
    if not headers:
        frappe.throw(_("File rỗng hoặc không có dòng header fieldname"))

    fieldname_map = _fieldname_map(doctype)
    unmapped = [h for h in headers if h and h.strip().lower() not in fieldname_map]

    errors: list[str] = []
    preview: list[dict] = []
    summary = {
        "total": len(rows), "created": 0, "updated": 0,
        "skipped": 0, "failed": 0, "would_create": 0, "would_update": 0,
    }
    is_submittable = bool(frappe.get_meta(doctype).is_submittable)

    is_dry = bool(cint(dry_run))

    for i, row in enumerate(rows, start=data_start):  # data_start = 2 hoặc 3
        row_errors_before = len(errors)
        doc_dict = _build_doc_dict(doctype, row, fieldname_map, errors, i)
        has_errors = len(errors) > row_errors_before

        name = doc_dict.get("name")
        exists = bool(name) and frappe.db.exists(doctype, name)
        if exists and not cint(update_existing):
            action = "skip"
        elif exists:
            action = "update"
        else:
            action = "create"

        if is_dry:
            preview_entry = {
                "line": i, "action": "error" if has_errors else action,
                "name": name, "fields_count": len([k for k in doc_dict.keys() if k != "doctype"]),
            }
            preview.append(preview_entry)
            if not has_errors:
                if action == "create":
                    summary["would_create"] += 1
                elif action == "update":
                    summary["would_update"] += 1
                else:
                    summary["skipped"] += 1
            else:
                summary["failed"] += 1
            continue

        # Commit mode — savepoint per row
        if has_errors:
            summary["failed"] += 1
            preview.append({"line": i, "action": "failed", "name": name,
                             "error": _("Lỗi validate")})
            continue

        sp = f"sc_import_{i}"
        try:
            frappe.db.savepoint(sp)
            if action == "skip":
                summary["skipped"] += 1
                preview.append({"line": i, "action": "skipped", "name": name,
                                 "reason": _("Đã tồn tại, không cập nhật")})
            elif action == "update":
                doc = frappe.get_doc(doctype, name)
                for k, v in doc_dict.items():
                    if k in ("doctype", "name") or k in SYSTEM_FIELDS:
                        continue
                    doc.set(k, v)
                doc.save()
                summary["updated"] += 1
                preview.append({"line": i, "action": "updated", "name": doc.name})
            else:
                # If user explicitly supplied name and doctype allows it, frappe uses it;
                # else autoname runs.
                doc = frappe.get_doc(doc_dict)
                doc.insert()
                if cint(submit_after) and is_submittable:
                    doc.submit()
                summary["created"] += 1
                preview.append({"line": i, "action": "created", "name": doc.name})
        except Exception as e:
            try:
                frappe.db.rollback(save_point=sp)
            except Exception:
                pass
            summary["failed"] += 1
            err_msg = str(e)[:400]
            errors.append(_("Dòng {0}: {1}").format(i, err_msg))
            preview.append({"line": i, "action": "failed", "name": name, "error": err_msg})

    if not is_dry:
        # Commit OK rows; failed ones already rolled back via savepoint
        frappe.db.commit()

    return {
        "dry_run": is_dry,
        "summary": summary,
        "preview": preview,
        "errors": errors,
        "unmapped_headers": unmapped,
        "headers": headers,
    }


# ---------------------------------------------------------------------------
# Convenience: list importable doctypes (server-side gatekeeper)
# ---------------------------------------------------------------------------

# Ledger doctypes — chỉ cho phép EXPORT (audit), không cho INSERT/UPDATE qua
# import để tránh hỏng tồn kho/sổ kế toán. SLE/GL Entry phải sinh qua voucher.
EXPORT_ONLY = {"SC Stock Ledger Entry", "SC GL Entry"}

# Registry mỗi module → doctypes hỗ trợ import/export. Đồng bộ với
# frontend/modules.js DT registry. Doctype không có ở đây vẫn dùng được nếu
# user có quyền (linh hoạt cho admin).
IMPORTABLE_BY_MODULE: dict[str, list[str]] = {
    "M0": ["SC Item", "SC Item Group", "SC UOM", "SC Supplier", "SC Warehouse",
            "Bin Location", "SC Department",
            "SC GL Account"],
    "M1": ["Framework Contract"],
    "M2": ["SC Material Request", "SC Purchase Order"],
    "M3": ["SC Purchase Receipt", "SC Quality Inspection"],
    "M4": ["SC Warehouse", "Bin Location", "SC Batch", "SC Stock Ledger Entry"],
    "M5": ["SC Batch"],
    "M6": ["SC Transfer Request", "SC Stock Entry"],
    "M8": ["SC Purchase Invoice", "SC Payment Entry", "SC GL Entry"],
    "M9": ["SC Inventory Count Sheet", "SC Stock Reconciliation"],
    "M10": ["SC Recall Notice", "SC Investigation Report"],
    "M11": ["SC Alert", "SC Alert Rule"],
}


@frappe.whitelist()
def list_importable() -> dict:
    """Liệt kê doctype theo module + permission flags cho UI."""
    out: dict[str, list[dict]] = {}
    for mod, doctypes in IMPORTABLE_BY_MODULE.items():
        out[mod] = []
        for dt in doctypes:
            if not frappe.db.exists("DocType", dt):
                continue
            can_read = frappe.has_permission(dt, "read")
            if not can_read:
                continue
            export_only = dt in EXPORT_ONLY
            out[mod].append({
                "doctype": dt,
                "label": frappe.get_meta(dt).get("label") or dt,
                "can_create": 0 if export_only else int(frappe.has_permission(dt, "create")),
                "can_write": 0 if export_only else int(frappe.has_permission(dt, "write")),
                "can_submit": int(frappe.has_permission(dt, "submit") or 0),
                "is_submittable": int(frappe.get_meta(dt).is_submittable or 0),
                "export_only": int(export_only),
            })
    return out
